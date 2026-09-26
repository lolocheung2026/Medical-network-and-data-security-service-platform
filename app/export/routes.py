"""报告导出模块 v0.3：推演报告 / 分类结果 → DOCX。

样式：微软雅黑正文 + 蓝色标题（乐叔文档偏好）。
数据来源：SimulationRun.result_json（推演）/ 已脱密字典 + 规则引擎（分类）。
"""
import io
import json
import time

from docx import Document
from docx.shared import Pt, RGBColor
from flask import Blueprint, send_file, session

from ..auth.routes import login_required
from ..classify.rules import LEVELS, classify_data_dict_csv
from ..models import SimulationRun, Upload

bp = Blueprint("export", __name__)

FONT_BODY = "微软雅黑"
COLOR_TITLE = RGBColor(0x18, 0x5F, 0xA5)


def _new_doc():
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = FONT_BODY
    style.font.size = Pt(10.5)
    style.element.rPr.rFonts.set(
        "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}eastAsia",
        FONT_BODY)
    return doc


def _title(doc, text, size=16):
    p = doc.add_paragraph()
    run = p.add_run(text)
    run.font.name = FONT_BODY
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = COLOR_TITLE
    return p


def _kv_table(doc, pairs):
    t = doc.add_table(rows=len(pairs), cols=2)
    t.style = "Table Grid"
    for i, (k, v) in enumerate(pairs):
        t.cell(i, 0).text = str(k)
        t.cell(i, 1).text = str(v)
    return t


def _docx_bytes(doc):
    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


@bp.route("/simulation/<int:run_id>")
@login_required
def simulation_export(run_id):
    rec = SimulationRun.query.filter_by(
        id=run_id, tenant_id=session["tenant_id"]).first_or_404()
    report = json.loads(rec.result_json)

    doc = _new_doc()
    _title(doc, "攻防推演报告")
    doc.add_paragraph(f"场景：{report.get('scenario')}（{report.get('scenario_id')}）")
    doc.add_paragraph(f"推演时间：{report.get('start_time')}")
    doc.add_paragraph(f"运行编号：{rec.id}")

    s = report.get("summary", {})
    _title(doc, "一、推演概要", size=13)
    _kv_table(doc, [
        ("攻击步总数", s.get("total_steps", "—")),
        ("攻击达成", s.get("completed", "—")),
        ("被拦截", s.get("blocked", "—")),
        ("节点攻陷", f"{s.get('nodes_compromised', '—')}/{s.get('total_nodes', '—')}"),
        ("防御生效", f"{s.get('defenses_active', '—')}/{s.get('total_defenses', '—')}"),
    ])

    _title(doc, "二、攻击步骤结果", size=13)
    for st in report.get("attack_steps", []):
        state = {"completed": "✓ 达成", "blocked": "⊘ 拦截"}.get(
            st.get("state"), f"· {st.get('state', '未知')}")
        doc.add_paragraph(f"{state}　{st.get('name')}（{st.get('mitre', '—')}）",
                          style="List Bullet")

    _title(doc, "三、整改建议", size=13)
    for m in report.get("remediation", []):
        doc.add_paragraph(
            f"{m.get('item')}（{m.get('category')} · 成本 {m.get('cost')}）",
            style="List Bullet")
        doc.add_paragraph(m.get("detail", ""))

    _title(doc, "四、结论边界声明", size=13)
    doc.add_paragraph(report.get("disclaimer") or
                      "本推演为纯仿真结果，不构成安全评估结论。")

    return send_file(
        _docx_bytes(doc), as_attachment=True,
        download_name=f"推演报告_{report.get('scenario_id')}_{rec.id}.docx",
        mimetype="application/vnd.openxmlformats-officedocument"
                 ".wordprocessingml.document")


@bp.route("/classify/<int:upload_id>")
@login_required
def classify_export(upload_id):
    rec = Upload.query.filter_by(
        id=upload_id, tenant_id=session["tenant_id"],
        kind=Upload.KIND_DATA_DICT, status=Upload.STATUS_DEIDENTIFIED,
    ).first_or_404()
    with open(rec.stored_path, encoding="utf-8") as fp:
        results = classify_data_dict_csv(fp.read())
    dist = {i: sum(1 for r in results if r["level"] == i) for i in range(1, 6)}

    doc = _new_doc()
    _title(doc, "数据分类分级结果")
    doc.add_paragraph(f"数据字典：{rec.original_filename}")
    doc.add_paragraph(f"生成时间：{time.strftime('%Y-%m-%d %H:%M')}")
    doc.add_paragraph(f"定级依据：GB/T 39725《健康医疗数据安全指南》")

    _title(doc, "一、敏感级别分布", size=13)
    _kv_table(doc, [(LEVELS[i], f"{dist[i]} 个字段") for i in range(1, 6)])

    _title(doc, "二、逐字段定级", size=13)
    t = doc.add_table(rows=1, cols=6)
    t.style = "Table Grid"
    for j, h in enumerate(["表名", "字段", "注释", "级别", "类别", "命中关键词"]):
        t.cell(0, j).text = h
    for r in results:
        row = t.add_row()
        for j, v in enumerate([r["table"], r["field"], r["comment"],
                               r["level_label"], r["category"],
                               r["hit_keyword"] or "—"]):
            row.cells[j].text = str(v)

    doc.add_paragraph()
    doc.add_paragraph("定级结果基于关键词规则引擎，供内部治理参考，"
                      "不构成等级保护测评结论。")

    return send_file(
        _docx_bytes(doc), as_attachment=True,
        download_name=f"分类分级_{rec.original_filename.rsplit('.', 1)[0]}.docx",
        mimetype="application/vnd.openxmlformats-officedocument"
                 ".wordprocessingml.document")
